import io
import re
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from amal.workbook_builder import (
    BOLD_FONT,
    CENTER,
    LEFT,
    PURPLE_FILL,
    RIGHT,
    SUPPLIER_TEXT,
    THIN_BORDER,
    THIN_SIDE,
    TOP_LEFT,
    WHITE_BOLD_FONT,
    apply_border_to_range,
    apply_outer_border_to_range,
    build_pack_list_sheet,
    compute_address_rows,
    count_actual_lines,
    fill_pack_list_items,
    fill_pack_list_sheet,
    get_layout,
    set_merged_block_row_heights,
    style_range,
    to_number_if_possible,
)


def build_comm_inv_static(worksheet) -> None:
    """Write Lenovo-specific commercial invoice layout with the extra Lenovo Qty column."""
    worksheet.title = "comm-inv"

    widths = {"A": 21, "B": 42, "C": 16, "D": 11, "E": 22, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18, "K": 10}
    for col, w in widths.items():
        worksheet.column_dimensions[col].width = w

    worksheet.row_dimensions[2].height = 28
    worksheet.merge_cells("A2:I2")
    worksheet["A2"] = "Commercial Invoice"
    worksheet["A2"].font = Font(bold=True, size=16)
    worksheet["A2"].alignment = CENTER

    worksheet.row_dimensions[4].height = 18
    worksheet.merge_cells("A4:I4")
    style_range(worksheet, "A4:I4", fill=PURPLE_FILL, border=THIN_BORDER)

    for row in (5, 6, 7):
        worksheet.row_dimensions[row].height = 18

    worksheet.merge_cells("A5:F5")
    worksheet.merge_cells("A6:F6")
    worksheet.merge_cells("A7:F7")
    worksheet.merge_cells("G5:I5")
    worksheet.merge_cells("G6:I6")
    worksheet.merge_cells("G7:I7")

    for addr in ("A5", "A6", "A7"):
        worksheet[addr].font = BOLD_FONT
        worksheet[addr].alignment = LEFT
    for addr in ("G5", "G6", "G7"):
        worksheet[addr].font = BOLD_FONT
        worksheet[addr].alignment = LEFT

    worksheet.row_dimensions[8].height = 18
    worksheet.merge_cells("A8:D8")
    worksheet.merge_cells("E8:F8")
    worksheet.merge_cells("G8:I8")
    worksheet["A8"] = "Bill To"
    worksheet["E8"] = "Ship To"
    worksheet["G8"] = "Supplier"
    style_range(worksheet, "A8:I8", fill=PURPLE_FILL, font=WHITE_BOLD_FONT, border=THIN_BORDER)
    worksheet["A8"].alignment = LEFT
    worksheet["E8"].alignment = LEFT
    worksheet["G8"].alignment = LEFT


def fill_comm_inv_sheet(worksheet, fields: dict, item_count: int) -> None:
    bill_to = fields.get("bill_to", "")
    ship_to = fields.get("ship_to", "")
    address_rows = compute_address_rows(bill_to, ship_to)
    L = get_layout(address_rows, item_count)

    worksheet["A5"] = ""
    worksheet["A6"] = f"Inco Terms: {fields.get('inco_terms', '')}"
    worksheet["A7"] = f"Customer PO: {fields.get('customer_po', '')}"
    for addr in ("A5", "A6", "A7"):
        worksheet[addr].font = BOLD_FONT
        worksheet[addr].alignment = LEFT

    worksheet["G5"] = f"Commercial Invoice No : {fields.get('commercial_invoice_no', '')}"
    worksheet["G6"] = f"Date: {fields.get('date', '')}"
    worksheet["G7"] = f"Currency: {fields.get('currency', '')}"
    for addr in ("G5", "G6", "G7"):
        worksheet[addr].font = BOLD_FONT
        worksheet[addr].alignment = LEFT

    addr_s, addr_e = L["addr_start"], L["addr_end"]
    worksheet.merge_cells(start_row=addr_s, start_column=1, end_row=addr_e, end_column=4)
    worksheet.merge_cells(start_row=addr_s, start_column=5, end_row=addr_e, end_column=6)
    worksheet.merge_cells(start_row=addr_s, start_column=7, end_row=addr_e, end_column=9)
    apply_outer_border_to_range(worksheet, addr_s, addr_e, 1, 9)

    c = worksheet.cell(row=addr_s, column=1)
    c.value = bill_to
    c.alignment = TOP_LEFT

    c = worksheet.cell(row=addr_s, column=5)
    c.value = ship_to
    c.alignment = TOP_LEFT

    c = worksheet.cell(row=addr_s, column=7)
    c.value = SUPPLIER_TEXT
    c.font = BOLD_FONT
    c.alignment = TOP_LEFT

    content_lines = max(
        count_actual_lines(bill_to),
        count_actual_lines(ship_to),
        count_actual_lines(SUPPLIER_TEXT),
    )
    set_merged_block_row_heights(worksheet, addr_s, addr_e, content_lines)

    hdr = L["header_row"]
    worksheet.row_dimensions[hdr].height = 22
    headers = ["Item Code", "Desc", "Case#", "Origin", "HS Code", "Qty", "Unit Price", "Amount", "Lenovo Qty"]
    for idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=hdr, column=idx)
        cell.value = header
        cell.fill = PURPLE_FILL
        cell.font = WHITE_BOLD_FONT
        cell.border = THIN_BORDER
        cell.alignment = LEFT if header in {"Item Code", "Desc"} else CENTER

    fr = L["freight_row"]
    worksheet.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=6)
    apply_border_to_range(worksheet, fr, fr, 7, 9)
    worksheet.cell(row=fr, column=7).value = "Freight Charges"
    worksheet.cell(row=fr, column=7).font = BOLD_FONT
    worksheet.cell(row=fr, column=7).alignment = RIGHT
    worksheet.cell(row=fr, column=7).border = THIN_BORDER
    worksheet.cell(row=fr, column=8).alignment = RIGHT
    worksheet.cell(row=fr, column=8).border = THIN_BORDER
    if fields.get("freight_charges", "") != "":
        worksheet.cell(row=fr, column=8).value = to_number_if_possible(fields.get("freight_charges", ""))

    tr = L["total_row"]
    worksheet.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=6)
    apply_border_to_range(worksheet, tr, tr, 7, 9)
    worksheet.cell(row=tr, column=7).value = "Total Amount"
    worksheet.cell(row=tr, column=7).font = BOLD_FONT
    worksheet.cell(row=tr, column=7).alignment = RIGHT
    worksheet.cell(row=tr, column=7).border = THIN_BORDER
    worksheet.cell(row=tr, column=8).font = BOLD_FONT
    worksheet.cell(row=tr, column=8).alignment = RIGHT
    worksheet.cell(row=tr, column=8).border = THIN_BORDER
    worksheet.cell(row=tr, column=8).value = f"=SUM(H{L['items_start']}:H{L['items_end']})+H{fr}"


def fill_comm_inv_items(worksheet, items: list[dict], address_rows: int) -> None:
    normal_items = [item for item in items if not item.get("append_only")]
    append_rows = [item for item in items if item.get("append_only")]
    L = get_layout(address_rows, len(normal_items))
    items_start = L["items_start"]

    max_desc_len = len("Desc")

    for offset, item in enumerate(normal_items):
        row = items_start + offset
        desc_val = str(item.get("desc", "")).replace("\n", " ").strip()
        max_desc_len = max(max_desc_len, len(desc_val))
        worksheet.row_dimensions[row].height = 15

        def w(col, value, align):
            c = worksheet.cell(row=row, column=col)
            c.value = value
            c.alignment = align
            c.border = THIN_BORDER

        w(1, item.get("item_code", ""), LEFT)
        w(2, desc_val, LEFT)
        w(3, item.get("case_no", ""), CENTER)
        w(4, item.get("origin", ""), CENTER)
        w(5, item.get("hs_code", ""), CENTER)
        w(6, item.get("qty", ""), CENTER)
        w(7, item.get("unit_price", ""), RIGHT)
        w(8, item.get("amount", ""), RIGHT)
        w(9, item.get("lenovo_qty", item.get("qty", "")), CENTER)

    worksheet.column_dimensions["B"].width = max_desc_len + 2
    apply_border_to_range(worksheet, L["items_start"], L["items_end"], 1, 9)


def _excel_engine_for_file(uploaded_file) -> str:
    suffix = Path(uploaded_file.name).suffix.lower()
    return "xlrd" if suffix == ".xls" else "openpyxl"


def _read_excel_sheet_rows(uploaded_file) -> pd.DataFrame:
    uploaded_file.seek(0)
    engine = _excel_engine_for_file(uploaded_file)
    df = pd.read_excel(uploaded_file, engine=engine, header=None)
    uploaded_file.seek(0)
    return df


def _extract_pack_list_table_rows(uploaded_file) -> list[list]:
    df = _read_excel_sheet_rows(uploaded_file)
    rows = df.fillna("").astype(str).values.tolist()

    start_index = 0
    for idx, row in enumerate(rows):
        text = " ".join(str(value).strip() for value in row if str(value).strip()).lower()
        if any(token in text for token in ("case no", "material", "part number", "qty", "gw(kg)", "description")):
            start_index = idx + 1
            break

    filtered_rows = []
    for row in rows[start_index:]:
        text = " ".join(str(value).strip() for value in row if str(value).strip())
        if not text or text.lower().startswith("total:"):
            continue
        filtered_rows.append(row)

    return filtered_rows


def fill_pack_list_sheet_static(worksheet, fields: dict, address_rows: int) -> None:
    """Static pack-list header layout for standalone Lenovo PL exports.

    This intentionally avoids formula references to another sheet, because the
    PL export is written as its own workbook in the ZIP output.
    """
    addr_s = 8
    addr_e = addr_s + address_rows - 1

    worksheet.merge_cells(start_row=addr_s, start_column=1, end_row=addr_e, end_column=4)
    worksheet.merge_cells(start_row=addr_s, start_column=5, end_row=addr_e, end_column=6)
    worksheet.merge_cells(start_row=addr_s, start_column=7, end_row=addr_e, end_column=8)
    apply_outer_border_to_range(worksheet, addr_s, addr_e, 1, 8)

    worksheet["G5"] = f"No. : {fields.get('commercial_invoice_no', '')}"
    worksheet["G6"] = f"Date : {fields.get('date', '')}"
    worksheet["G5"].font = BOLD_FONT
    worksheet["G6"].font = BOLD_FONT
    worksheet["G5"].alignment = LEFT
    worksheet["G6"].alignment = LEFT

    c = worksheet.cell(row=addr_s, column=1)
    c.value = fields.get("bill_to", "")
    c.alignment = TOP_LEFT

    c = worksheet.cell(row=addr_s, column=5)
    c.value = fields.get("ship_to", "")
    c.alignment = TOP_LEFT

    c = worksheet.cell(row=addr_s, column=7)
    c.value = SUPPLIER_TEXT
    c.font = BOLD_FONT
    c.alignment = TOP_LEFT

    set_merged_block_row_heights(
        worksheet,
        addr_s,
        addr_e,
        max(
            count_actual_lines(fields.get("bill_to", "")),
            count_actual_lines(fields.get("ship_to", "")),
            count_actual_lines(SUPPLIER_TEXT),
        ),
    )

    hdr_row = addr_e + 1
    worksheet.row_dimensions[hdr_row].height = 22
    headers = ["Item Code", "Desc", "Case#", "Origin", "HS Code", "Qty", "Weight", "Package"]
    for idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=hdr_row, column=idx)
        cell.value = header
        cell.fill = PURPLE_FILL
        cell.font = WHITE_BOLD_FONT
        cell.border = THIN_BORDER
        cell.alignment = LEFT if header in {"Item Code", "Desc"} else CENTER


def _copy_pack_list_sheet_with_header(
    workbook: Workbook,
    uploaded_file,
    sheet_name: str,
    pack_list_fields: dict,
    address_rows: int,
) -> None:
    worksheet = workbook.create_sheet(title=sheet_name)
    build_pack_list_sheet(worksheet)
    fill_pack_list_sheet(worksheet, pack_list_fields, address_rows)

    rows = _extract_pack_list_table_rows(uploaded_file)
    hdr_row = 8 + address_rows
    start_row = hdr_row + 2

    for row_idx, row_values in enumerate(rows, start=start_row):
        for col_idx, value in enumerate(row_values, start=1):
            worksheet.cell(row=row_idx, column=col_idx, value=value)


def create_workbook_bytes(
    comm_inv_fields: dict,
    comm_inv_items: list[dict],
    comm_inv_unmatched_items: list[dict],
    pack_list_fields: dict,
    pack_list_items: list[dict],
    comm_inv_df: pd.DataFrame,
    pack_list_df: pd.DataFrame,
    extra_pack_list_files: list | None = None,
    include_pack_list_sheet: bool = True,
) -> io.BytesIO:
    bill_to = comm_inv_fields.get("bill_to", "")
    ship_to = comm_inv_fields.get("ship_to", "")
    address_rows = compute_address_rows(bill_to, ship_to)

    workbook = Workbook()

    comm_sheet = workbook.active
    normal_items = [item for item in comm_inv_items if not item.get("append_only")]
    append_rows = [item for item in comm_inv_items if item.get("append_only")]

    build_comm_inv_static(comm_sheet)
    fill_comm_inv_items(comm_sheet, comm_inv_items, address_rows)
    fill_comm_inv_sheet(comm_sheet, comm_inv_fields, len(normal_items))

    if append_rows:
        append_start = comm_sheet.max_row + 1
        comm_sheet.cell(row=append_start, column=10, value="Part Number")
        comm_sheet.cell(row=append_start, column=11, value="Qty")
        for offset, item in enumerate(append_rows, start=1):
            row = append_start + offset
            comm_sheet.cell(row=row, column=10, value=item.get("item_code", ""))
            comm_sheet.cell(row=row, column=11, value=item.get("qty", ""))
            comm_sheet.row_dimensions[row].height = 15

    if include_pack_list_sheet:
        if extra_pack_list_files:
            _copy_pack_list_sheet_with_header(
                workbook,
                extra_pack_list_files[0],
                "pack_list",
                pack_list_fields,
                address_rows,
            )
            for idx, extra_file in enumerate(extra_pack_list_files[1:], start=1):
                _copy_pack_list_sheet_with_header(
                    workbook,
                    extra_file,
                    f"pack_list_{idx}",
                    pack_list_fields,
                    address_rows,
                )
        else:
            pack_sheet = workbook.create_sheet("pack_list")
            build_pack_list_sheet(pack_sheet)
            fill_pack_list_sheet(pack_sheet, pack_list_fields, address_rows)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
